import os
from typing import List, Dict, Any, Optional, Tuple
from database import Database, db as default_db
from models.subject import SubjectBlockModel
from services.abbreviations import get_acronym, get_legend_mapping, is_subject_applicable_to_tipo

# ReportLab imports for PDF
from reportlab.lib.pagesizes import letter, A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

# Excel import
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Word import
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, logo_path: Optional[str] = None, watermark_text: Optional[str] = None, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []
        self.logo_path = logo_path
        self.watermark_text = watermark_text

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        self.saveState()
        pw, ph = self._pagesize
        
        # Marca d'água em diagonal (se configurada)
        if self.watermark_text:
            self.setFont("Helvetica-Bold", 42)
            self.setFillColor(colors.HexColor("#242D64"))
            self.setFillAlpha(0.08)
            self.translate(pw / 2.0, ph / 2.0)
            self.rotate(45)
            self.drawCentredString(0, 0, self.watermark_text)
            self.restoreState()
            self.saveState()

        # Cabeçalho decorativo
        self.setStrokeColor(colors.HexColor("#00A9A4"))
        self.setLineWidth(1.5)
        self.line(1.0 * cm, ph - 1.2 * cm, pw - 1.0 * cm, ph - 1.2 * cm)

        # Logotipo no cabeçalho se existir
        logo_to_use = self.logo_path
        if not logo_to_use or not os.path.exists(logo_to_use):
            if os.path.exists("logo.png"):
                logo_to_use = "logo.png"

        if logo_to_use and os.path.exists(logo_to_use):
            try:
                self.drawImage(logo_to_use, pw - 3.8 * cm, ph - 1.1 * cm, width=2.6 * cm, height=0.9 * cm, preserveAspectRatio=True, mask='auto')
            except Exception:
                pass

        # Rodapé
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(1.0 * cm, 1.4 * cm, pw - 1.0 * cm, 1.4 * cm)

        self.drawString(1.0 * cm, 0.9 * cm, "Corretor de Provas - Relatório Oficial")
        self.drawRightString(pw - 1.0 * cm, 0.9 * cm, f"Página {self._pageNumber} de {page_count}")
        self.restoreState()


class ReportExporter:
    def __init__(self, db: Database = default_db):
        self.db = db
        self.subject_model = SubjectBlockModel(db)

    def _get_branding(self) -> Tuple[Optional[str], Optional[str]]:
        logo_path = self.db.get_config("logo_path")
        if not logo_path or not os.path.exists(logo_path):
            if os.path.exists("logo.png"):
                logo_path = "logo.png"
        watermark = self.db.get_config("watermark_text", "CORRETOR DE PROVAS")
        return logo_path, watermark

    def _get_block_mapping(self, exam: Dict[str, Any]) -> Tuple[Dict[str, str], List[str]]:
        disc_db = {d["nome"].strip().lower(): d.get("bloco_nome") for d in self.subject_model.list_subjects()}
        blocos_list = exam.get("blocos_list", [])
        exam_block_names = [b["nome"] for b in blocos_list] if blocos_list else []
        return disc_db, exam_block_names

    # --- GERAR PDF DA PROVA (DISCIPLINA EM LANDSCAPE / GERAL E BLOCO EM PORTRAIT) ---
    def export_exam_pdf(
        self, 
        exam: Dict[str, Any], 
        results: List[Dict[str, Any]], 
        output_path: str, 
        turma_subtitle: str = "Geral - Todas as Turmas",
        view_mode: str = "geral"
    ):
        logo_path, watermark = self._get_branding()
        possui_redacao = bool(exam.get("possui_redacao", False))
        
        is_landscape = (view_mode == "disciplina")
        pagesize = landscape(A4) if is_landscape else A4

        doc = SimpleDocTemplate(
            output_path,
            pagesize=pagesize,
            leftMargin=1.2 * cm if is_landscape else 1.0 * cm,
            rightMargin=1.2 * cm if is_landscape else 1.0 * cm,
            topMargin=1.6 * cm,
            bottomMargin=1.8 * cm
        )

        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=17,
            textColor=colors.HexColor("#242D64"),
            alignment=0
        )
        
        subtitle_style = ParagraphStyle(
            "CustomSubTitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#64748B")
        )

        table_header_style = ParagraphStyle(
            "TableHeader",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7.5,
            leading=9,
            textColor=colors.white,
            alignment=1
        )

        table_header_top_style = ParagraphStyle(
            "TableHeaderTop",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=10,
            textColor=colors.white,
            alignment=1
        )

        table_cell_style = ParagraphStyle(
            "TableCell",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9,
            textColor=colors.HexColor("#1E293B"),
            alignment=1
        )

        table_cell_student_style = ParagraphStyle(
            "TableCellStudent",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9,
            textColor=colors.HexColor("#1E293B"),
            alignment=0
        )

        legend_style = ParagraphStyle(
            "LegendStyle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=10,
            textColor=colors.HexColor("#475569")
        )

        story = []

        # Título e Subtítulo por Turma
        mode_label = view_mode.upper()
        story.append(Paragraph(f"Relatório de Resultados - {exam.get('nome', 'Prova')} (Modo: {mode_label})", title_style))
        story.append(Paragraph(f"<b>Turma:</b> {turma_subtitle} | Data: {exam.get('data', 'N/A')} | Bloco(s): {exam.get('bloco_nome') or 'Geral'} | Valor Total: {exam.get('valor_total', 10.0)} pts | Alunos Processados: {len(results)}", subtitle_style))

        disc_db_map, exam_blocks = self._get_block_mapping(exam)
        mapped_subjects = exam.get("layout_config", {}).get("subjects", [])
        exam_disc_list = []
        for s in mapped_subjects:
            n = s.get("nome", "").strip()
            if n and n.lower() != "geral" and n not in exam_disc_list:
                exam_disc_list.append(n)

        if not exam_disc_list:
            disc_set = set()
            for r in results:
                det = r.get("detalhes_disciplinas", {})
                for d_name in det.keys():
                    if d_name.strip().lower() != "geral":
                        disc_set.add(d_name.strip())
            exam_disc_list = sorted(list(disc_set))

        if view_mode in ["geral", "bloco"]:
            block_set = set(exam_blocks) if exam_blocks else {disc_db_map.get(d.lower()) or "Geral" for d in exam_disc_list}
            mid_full = sorted(list(block_set))
            group_title = "Acertos" if view_mode == "geral" else "Notas"
        else:
            mid_full = exam_disc_list
            group_title = "Notas"

        # Legenda antes da tabela
        leg_map = get_legend_mapping(mid_full)
        if leg_map:
            leg_items = [f"<b>{k}</b>: {v}" for k, v in leg_map.items()]
            leg_str = "  |  ".join(leg_items)
            story.append(Spacer(1, 4))
            story.append(Paragraph(f"<b>Legenda:</b> {leg_str}", legend_style))

        story.append(Spacer(1, 10))

        mid_siglas = [get_acronym(m) for m in mid_full]
        num_mid = len(mid_siglas)

        # Montar colunas finais do cabeçalho
        right_headers = []
        if view_mode == "geral":
            right_headers.append("Total")
        if possui_redacao:
            red_title = "RED" if view_mode == "disciplina" else "Redação"
            right_headers.append(red_title)
        if possui_redacao or view_mode != "geral":
            right_headers.append("Nota Final")

        total_cols = 5 + num_mid + len(right_headers)
        row1 = [Paragraph(f"<b>{group_title}</b>", table_header_top_style)] + [Paragraph("", table_header_style) for _ in range(total_cols - 1)]

        row2 = [
            Paragraph("", table_header_style),
            Paragraph("Matrícula", table_header_style),
            Paragraph("Aluno", table_header_style),
            Paragraph("Turma", table_header_style),
            Paragraph("Tipo", table_header_style),
        ]
        for sig in mid_siglas:
            row2.append(Paragraph(sig, table_header_style))

        for rh in right_headers:
            row2.append(Paragraph(rh, table_header_style))

        table_data = [row1, row2]

        for row_idx, r in enumerate(results, 1):
            det = r.get("detalhes_disciplinas", {})
            aluno_nome = str(r.get("aluno_nome", "Aluno Não Cadastrado")).replace("\n", " ").replace("\r", "")
            row = [
                Paragraph(str(row_idx), table_cell_style),
                Paragraph(str(r.get("aluno_matricula", "")), table_cell_style),
                Paragraph(f"<nobr>{aluno_nome}</nobr>", table_cell_student_style),
                Paragraph(str(r.get("aluno_turma", "N/A")), table_cell_style),
                Paragraph(str(r.get("tipo_prova", "")), table_cell_style),
            ]

            if view_mode == "geral":
                for b_name in mid_full:
                    b_ac = sum(d_v.get("acertos", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    row.append(Paragraph(str(b_ac), table_cell_style))
                row.append(Paragraph(f"{r.get('total_acertos', 0)}/{r.get('total_questoes', 0)}", table_cell_style))

            elif view_mode == "bloco":
                for b_name in mid_full:
                    b_ac = sum(d_v.get("acertos", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    b_tot = sum(d_v.get("total", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    b_nota = (b_ac / b_tot * 10.0) if b_tot > 0 else 0.0
                    row.append(Paragraph(f"{b_nota:.2f}", table_cell_style))

            else: # view_mode == "disciplina"
                for d_name in mid_full:
                    if not is_subject_applicable_to_tipo(d_name, r.get("tipo_prova", "")):
                        row.append(Paragraph("-", table_cell_style))
                    else:
                        d_info = det.get(d_name)
                        if not d_info:
                            for k, v in det.items():
                                if k.strip().lower() == d_name.strip().lower():
                                    d_info = v
                                    break
                        if not d_info:
                            d_info = {}
                        tot = d_info.get("total", 0)
                        ac = d_info.get("acertos", 0)
                        n_disc = (ac / tot * 10.0) if tot > 0 else d_info.get("nota", 0.0)
                        row.append(Paragraph(f"{n_disc:.2f}", table_cell_style))

            if possui_redacao:
                n_red = r.get("nota_redacao")
                n_red_str = f"{n_red:.2f}" if n_red is not None else "0.00"
                row.append(Paragraph(n_red_str, table_cell_style))

            if possui_redacao or view_mode != "geral":
                n_red_val = r.get("nota_redacao") or 0.0 if possui_redacao else 0.0
                fn_val = float(r.get("total_acertos", 0)) + n_red_val
                row.append(Paragraph(f"<b>{fn_val:.2f}</b>", table_cell_style))

            table_data.append(row)

        # Ajuste de largura de colunas (Landscape = 27.3cm / Portrait = 19.0cm)
        total_page_w = 27.3 * cm if is_landscape else 19.0 * cm
        max_name_len = max([len(str(r.get("aluno_nome", ""))) for r in results], default=15)
        max_name_w = 7.5 * cm if is_landscape else 5.5 * cm
        aluno_col_w = max(3.2 * cm, min(max_name_w, (max_name_len * 0.16 + 0.6) * cm))

        right_cols_cnt = len(right_headers)
        fixed_w = 0.8 * cm + 1.8 * cm + aluno_col_w + 1.6 * cm + 1.2 * cm + (1.6 * cm * right_cols_cnt)
        avail_space = total_page_w - fixed_w
        mid_width = (avail_space / num_mid) if num_mid > 0 else 1.2 * cm
        if mid_width < 1.0 * cm:
            mid_width = 1.0 * cm

        col_widths = [0.8 * cm, 1.8 * cm, aluno_col_w, 1.6 * cm, 1.2 * cm] + [mid_width] * num_mid + [1.6 * cm] * right_cols_cnt

        t_style = [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor("#242D64")),
            ('TEXTCOLOR', (0, 0), (-1, 1), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('SPAN', (0, 0), (-1, 0)),
            ('ROWBACKGROUNDS', (0, 2), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
        ]

        t = Table(table_data, colWidths=col_widths, repeatRows=2)
        t.setStyle(TableStyle(t_style))

        story.append(t)

        def make_canvas(*args, **kwargs):
            return NumberedCanvas(*args, logo_path=logo_path, watermark_text=watermark, **kwargs)

        doc.build(story, canvasmaker=make_canvas)

    # --- GERAR PDF INDIVIDUAL DO ALUNO ---
    def export_individual_pdf(self, student_info: Dict[str, Any], exam_info: Dict[str, Any], result: Dict[str, Any], output_path: str):
        logo_path, watermark = self._get_branding()

        doc = SimpleDocTemplate(
            output_path,
            pagesize=A4,
            leftMargin=1.2 * cm,
            rightMargin=1.2 * cm,
            topMargin=1.8 * cm,
            bottomMargin=2.0 * cm
        )

        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            "IndTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=18,
            textColor=colors.HexColor("#242D64"),
            alignment=0
        )
        
        body_style = ParagraphStyle(
            "IndBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#1E293B")
        )

        header_style = ParagraphStyle(
            "IndHeader",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            textColor=colors.white,
            alignment=1
        )

        story = []

        # Cabeçalho do Boletim
        e_nome = exam_info.get('nome', 'Prova') if exam_info else "Histórico do Aluno"
        story.append(Paragraph(f"Boletim Individual - {e_nome}", title_style))
        story.append(Spacer(1, 10))

        # Card do Aluno
        aluno_nome = str(student_info.get('nome', 'N/A')).replace("\n", " ").replace("\r", "")
        n_final = result.get('nota_final', 0.0) if result else 0.0
        v_tot = exam_info.get('valor_total', 10.0) if exam_info else 10.0
        pct_ac = result.get('percentual_acertos', 0.0) if result else 0.0
        tot_ac = result.get('total_acertos', 0) if result else 0
        tot_q = result.get('total_questoes', 0) if result else 0

        info_data = [
            [
                Paragraph(f"<b>Aluno:</b> <nobr>{aluno_nome}</nobr>", body_style),
                Paragraph(f"<b>Matrícula:</b> {student_info.get('matricula', result.get('aluno_matricula', 'N/A'))}", body_style)
            ],
            [
                Paragraph(f"<b>Turma:</b> {student_info.get('turma', 'N/A')}", body_style),
                Paragraph(f"<b>Tipo da Prova:</b> {result.get('tipo_prova', 'N/A') if result else 'N/A'}", body_style)
            ],
            [
                Paragraph(f"<b>Nota Final da Prova:</b> <font size=11 color='#00A9A4'><b>{n_final:.2f}</b> / {v_tot}</font>", body_style),
                Paragraph(f"<b>Aproveitamento Geral:</b> {pct_ac:.1f}% ({tot_ac}/{tot_q} acertos)", body_style)
            ]
        ]
        info_table = Table(info_data, colWidths=[9.3 * cm, 9.3 * cm])
        info_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(info_table)
        story.append(Spacer(1, 12))

        # Detalhamento por Disciplina
        story.append(Paragraph("<b>Desempenho por Disciplina (Nota de 0 a 10 por disciplina)</b>", title_style))
        story.append(Spacer(1, 6))

        disc_headers = ["Disciplina", "Acertos", "Total Questões", "% Acertos", "Nota Disciplina"]
        disc_table_data = [[Paragraph(h, header_style) for h in disc_headers]]

        detalhes = result.get("detalhes_disciplinas", {}) if result else {}
        for disc_name, disc_data in detalhes.items():
            if disc_name.strip().lower() == "geral":
                continue
            ac = disc_data.get("acertos", 0)
            tot = disc_data.get("total", 0)
            pct = disc_data.get("percentual", 0.0)
            n_disc = (ac / tot * 10.0) if tot > 0 else disc_data.get("nota", 0.0)
            disc_table_data.append([
                Paragraph(disc_name, body_style),
                Paragraph(str(ac), body_style),
                Paragraph(str(tot), body_style),
                Paragraph(f"{pct:.1f}%", body_style),
                Paragraph(f"<b>{n_disc:.2f}</b>", body_style)
            ])

        if len(disc_table_data) > 1:
            dt = Table(disc_table_data, colWidths=[6.6 * cm, 2.5 * cm, 2.5 * cm, 3 * cm, 4 * cm])
            dt.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#242D64")),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('ALIGN', (0, 1), (0, -1), 'LEFT'),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ('PADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(dt)

        def make_canvas(*args, **kwargs):
            return NumberedCanvas(*args, logo_path=logo_path, watermark_text=watermark, **kwargs)

        doc.build(story, canvasmaker=make_canvas)

    # --- GERAR EXCEL INDIVIDUAL DO ALUNO ---
    def export_individual_excel(self, student_info: Dict[str, Any], stats_or_result: Dict[str, Any], output_path: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Boletim Aluno"
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT

        header_fill = PatternFill(start_color="242D64", end_color="242D64", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="242D64")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
        )

        aluno_nome = str(student_info.get('nome', '')).replace("\n", " ").replace("\r", "")
        ws.append([f"Boletim Individual do Aluno: {aluno_nome}"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Matrícula: {student_info.get('matricula', '')} | Turma: {student_info.get('turma', '')}"])
        ws.cell(row=2, column=1).font = subtitle_font
        ws.append([])

        history = stats_or_result.get("historico", [])
        if not history and "nota_final" in stats_or_result:
            history = [stats_or_result]

        headers = ["Data Prova", "Nome da Prova", "Tipo", "Acertos", "% Acertos", "Nota Final", "Status"]
        ws.append(headers)
        h_row = 4
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=h_row, column=c_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for h in history:
            row_data = [
                h.get("prova_data", "N/A"),
                h.get("prova_nome", "Prova"),
                h.get("tipo_prova", "N/A"),
                f"{h.get('total_acertos', 0)}/{h.get('total_questoes', 0)}",
                f"{h.get('percentual_acertos', 0.0):.1f}%",
                h.get("nota_final", 0.0),
                h.get("status_presenca", "Presente")
            ]
            ws.append(row_data)
            curr = ws.max_row
            for c_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=curr, column=c_idx)
                c.border = thin_border
                c.alignment = center_align if c_idx in [1, 3, 4, 5, 6, 7] else left_align

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(output_path)

    # --- GERAR WORD INDIVIDUAL DO ALUNO ---
    def export_individual_word(self, student_info: Dict[str, Any], stats_or_result: Dict[str, Any], output_path: str):
        doc = docx.Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT

        aluno_nome = str(student_info.get('nome', '')).replace("\n", " ").replace("\r", "")
        h1 = doc.add_heading(level=1)
        run = h1.add_run(f"Boletim Individual: {aluno_nome}")
        run.font.color.rgb = RGBColor(0x24, 0x2D, 0x64)

        p = doc.add_paragraph()
        p.add_run(f"Matrícula: {student_info.get('matricula', '')} | Turma: {student_info.get('turma', '')}\n")

        history = stats_or_result.get("historico", [])
        if not history and "nota_final" in stats_or_result:
            history = [stats_or_result]

        headers = ["Data Prova", "Nome da Prova", "Tipo", "Acertos", "% Acertos", "Nota Final", "Status"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].font.bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x24, 0x2D, 0x64)

        for h in history:
            row_cells = table.add_row().cells
            row_cells[0].text = str(h.get("prova_data", "N/A"))
            row_cells[1].text = str(h.get("prova_nome", "N/A"))
            row_cells[2].text = str(h.get("tipo_prova", "N/A"))
            row_cells[3].text = f"{h.get('total_acertos', 0)}/{h.get('total_questoes', 0)}"
            row_cells[4].text = f"{h.get('percentual_acertos', 0.0):.1f}%"
            row_cells[5].text = f"{h.get('nota_final', 0.0):.2f}"
            row_cells[6].text = str(h.get("status_presenca", "Presente"))

        doc.save(output_path)

    # --- EXPORTAR PARA EXCEL (DISCIPLINA EM LANDSCAPE / OUTROS EM PORTRAIT) ---
    def export_exam_excel(
        self, 
        exam: Dict[str, Any], 
        results: List[Dict[str, Any]], 
        output_path: str, 
        turma_subtitle: str = "Geral - Todas as Turmas",
        view_mode: str = "geral"
    ):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resultados"
        if view_mode == "disciplina":
            ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        else:
            ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT

        possui_redacao = bool(exam.get("possui_redacao", False))

        header_fill = PatternFill(start_color="242D64", end_color="242D64", fill_type="solid")
        header_group_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="242D64")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
        )

        ws.append([f"Relatório de Prova: {exam.get('nome', '')} (Modo: {view_mode.upper()})"])
        ws.cell(row=1, column=1).font = title_font
        ws.append([f"Turma: {turma_subtitle} | Data: {exam.get('data', '')} | Bloco(s): {exam.get('bloco_nome', 'Geral')} | Total Alunos Processados: {len(results)} | Valor Total Prova: {exam.get('valor_total', 10.0)}"])
        ws.cell(row=2, column=1).font = subtitle_font

        disc_db_map, exam_blocks = self._get_block_mapping(exam)
        mapped_subjects = exam.get("layout_config", {}).get("subjects", [])
        exam_disc_list = []
        for s in mapped_subjects:
            n = s.get("nome", "").strip()
            if n and n.lower() != "geral" and n not in exam_disc_list:
                exam_disc_list.append(n)

        if not exam_disc_list:
            disc_set = set()
            for r in results:
                det = r.get("detalhes_disciplinas", {})
                for d in det.keys():
                    if d.strip().lower() != "geral":
                        disc_set.add(d.strip())
            exam_disc_list = sorted(list(disc_set))

        if view_mode in ["geral", "bloco"]:
            block_set = set(exam_blocks) if exam_blocks else {disc_db_map.get(d.lower()) or "Geral" for d in exam_disc_list}
            mid_full = sorted(list(block_set))
            group_title = "Acertos" if view_mode == "geral" else "Notas"
        else:
            mid_full = exam_disc_list
            group_title = "Notas"

        mid_siglas = [get_acronym(m) for m in mid_full]

        base_left = ["", "Matrícula", "Aluno", "Turma", "Tipo Prova"]
        base_right = []
        if view_mode == "geral":
            base_right.append("Total Acertos")
        if possui_redacao:
            red_title = "RED" if view_mode == "disciplina" else "Redação"
            base_right.append(red_title)
        if possui_redacao or view_mode != "geral":
            base_right.append("Nota Final")

        headers_l5 = base_left + mid_siglas + base_right
        total_cols_cnt = len(headers_l5)

        # Legenda antes da tabela mesclada na largura da tabela
        leg_map = get_legend_mapping(mid_full)
        if leg_map:
            leg_items = [f"{sig} = {full_name}" for sig, full_name in leg_map.items()]
            ws.append([f"Legenda: {' | '.join(leg_items)}"])
            r_leg = ws.max_row
            ws.merge_cells(start_row=r_leg, start_column=1, end_row=r_leg, end_column=total_cols_cnt)
            cell_leg = ws.cell(row=r_leg, column=1)
            cell_leg.font = Font(name="Calibri", size=9, bold=True, color="242D64")
            cell_leg.alignment = Alignment(horizontal="left", vertical="center")

        ws.append([])

        row_grp = ws.max_row + 1
        ws.append([""] * total_cols_cnt)
        row_hdr = ws.max_row + 1
        ws.append(headers_l5)

        ws.merge_cells(start_row=row_grp, start_column=1, end_row=row_grp, end_column=total_cols_cnt)
        cell_grp = ws.cell(row=row_grp, column=1)
        cell_grp.value = group_title.upper()
        cell_grp.fill = header_group_fill
        cell_grp.font = header_font
        cell_grp.alignment = center_align

        for col_idx in range(1, total_cols_cnt + 1):
            c = ws.cell(row=row_hdr, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = center_align

        for row_idx, r in enumerate(results, 1):
            det = r.get("detalhes_disciplinas", {})
            aluno_nome = str(r.get("aluno_nome", "")).replace("\n", " ").replace("\r", "")
            row_data = [
                row_idx,
                r.get("aluno_matricula", ""),
                aluno_nome,
                r.get("aluno_turma", ""),
                r.get("tipo_prova", ""),
            ]

            if view_mode == "geral":
                for b_name in mid_full:
                    b_ac = sum(d_v.get("acertos", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    row_data.append(b_ac)
                row_data.append(r.get("total_acertos", 0))

            elif view_mode == "bloco":
                for b_name in mid_full:
                    b_ac = sum(d_v.get("acertos", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    b_tot = sum(d_v.get("total", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    b_nota = (b_ac / b_tot * 10.0) if b_tot > 0 else 0.0
                    row_data.append(round(b_nota, 2))

            else: # view_mode == "disciplina"
                for d_name in mid_full:
                    if not is_subject_applicable_to_tipo(d_name, r.get("tipo_prova", "")):
                        row_data.append("-")
                    else:
                        d_info = det.get(d_name)
                        if not d_info:
                            for k, v in det.items():
                                if k.strip().lower() == d_name.strip().lower():
                                    d_info = v
                                    break
                        if not d_info:
                            d_info = {}
                        tot = d_info.get("total", 0)
                        ac = d_info.get("acertos", 0)
                        n_disc = (ac / tot * 10.0) if tot > 0 else d_info.get("nota", 0.0)
                        row_data.append(round(n_disc, 2))

            if possui_redacao:
                n_red = r.get("nota_redacao")
                row_data.append(round(n_red, 2) if n_red is not None else 0.0)

            if possui_redacao or view_mode != "geral":
                n_red_val = r.get("nota_redacao") or 0.0 if possui_redacao else 0.0
                fn_val = float(r.get("total_acertos", 0)) + n_red_val
                row_data.append(round(fn_val, 2))

            ws.append(row_data)
            current_row = ws.max_row
            
            for col_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=current_row, column=col_idx)
                c.border = thin_border
                if col_idx >= 6:
                    if view_mode != "geral" or col_idx != len(row_data):
                        c.number_format = "0.00"
                    c.alignment = center_align
                elif col_idx in [1, 2, 4, 5]:
                    c.alignment = center_align
                else:
                    c.alignment = left_align

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        wb.save(output_path)

    # --- EXPORTAR PARA WORD (DISCIPLINA EM LANDSCAPE / OUTROS EM PORTRAIT) ---
    def export_exam_word(
        self, 
        exam: Dict[str, Any], 
        results: List[Dict[str, Any]], 
        output_path: str, 
        turma_subtitle: str = "Geral - Todas as Turmas",
        view_mode: str = "geral"
    ):
        doc = docx.Document()
        section = doc.sections[0]
        if view_mode == "disciplina":
            section.orientation = WD_ORIENT.LANDSCAPE
            new_w, new_h = section.page_height, section.page_width
            section.page_width = new_w
            section.page_height = new_h
        else:
            section.orientation = WD_ORIENT.PORTRAIT

        possui_redacao = bool(exam.get("possui_redacao", False))

        h1 = doc.add_heading(level=1)
        run = h1.add_run(f"Relatório de Prova: {exam.get('nome', '')} (Modo: {view_mode.upper()})")
        run.font.color.rgb = RGBColor(0x24, 0x2D, 0x64)

        p = doc.add_paragraph()
        p.add_run(f"Turma: {turma_subtitle}\n")
        p.add_run(f"Data: {exam.get('data', '')} | Bloco(s): {exam.get('bloco_nome') or 'Geral'} | Valor Total Prova: {exam.get('valor_total', 10.0)} pts\n")
        p.add_run(f"Total de Alunos Processados: {len(results)}")

        disc_db_map, exam_blocks = self._get_block_mapping(exam)
        mapped_subjects = exam.get("layout_config", {}).get("subjects", [])
        exam_disc_list = []
        for s in mapped_subjects:
            n = s.get("nome", "").strip()
            if n and n.lower() != "geral" and n not in exam_disc_list:
                exam_disc_list.append(n)

        if not exam_disc_list:
            disc_set = set()
            for r in results:
                det = r.get("detalhes_disciplinas", {})
                for d in det.keys():
                    if d.strip().lower() != "geral":
                        disc_set.add(d.strip())
            exam_disc_list = sorted(list(disc_set))

        if view_mode in ["geral", "bloco"]:
            block_set = set(exam_blocks) if exam_blocks else {disc_db_map.get(d.lower()) or "Geral" for d in exam_disc_list}
            mid_full = sorted(list(block_set))
        else:
            mid_full = exam_disc_list

        leg_map = get_legend_mapping(mid_full)
        if leg_map:
            p_leg = doc.add_paragraph()
            r_head = p_leg.add_run("Legenda: ")
            r_head.bold = True
            r_head.font.color.rgb = RGBColor(0x24, 0x2D, 0x64)
            leg_str = " | ".join([f"{sig} = {full_name}" for sig, full_name in leg_map.items()])
            p_leg.add_run(leg_str)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        mid_siglas = [get_acronym(m) for m in mid_full]

        right_headers = []
        if view_mode == "geral":
            right_headers.append("Total Acertos")
        if possui_redacao:
            red_title = "RED" if view_mode == "disciplina" else "Redação"
            right_headers.append(red_title)
        if possui_redacao or view_mode != "geral":
            right_headers.append("Nota Final")

        headers = ["", "Matrícula", "Aluno", "Turma", "Tipo"] + mid_siglas + right_headers

        table = doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].font.bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0x24, 0x2D, 0x64)

        for row_idx, r in enumerate(results, 1):
            det = r.get("detalhes_disciplinas", {})
            aluno_nome = str(r.get("aluno_nome", "")).replace("\n", " ").replace("\r", "")
            row_cells = table.add_row().cells
            row_cells[0].text = str(row_idx)
            row_cells[1].text = str(r.get("aluno_matricula", ""))
            row_cells[2].text = aluno_nome
            row_cells[3].text = str(r.get("aluno_turma", ""))
            row_cells[4].text = str(r.get("tipo_prova", ""))

            col_idx = 5
            if view_mode == "geral":
                for b_name in mid_full:
                    b_ac = sum(d_v.get("acertos", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    row_cells[col_idx].text = str(b_ac)
                    col_idx += 1
                row_cells[col_idx].text = f"{r.get('total_acertos', 0)}/{r.get('total_questoes', 0)}"
                col_idx += 1

            elif view_mode == "bloco":
                for b_name in mid_full:
                    b_ac = sum(d_v.get("acertos", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    b_tot = sum(d_v.get("total", 0) for d_k, d_v in det.items() if d_k.strip().lower() != "geral" and (disc_db_map.get(d_k.strip().lower()) == b_name or len(mid_full) == 1))
                    b_nota = (b_ac / b_tot * 10.0) if b_tot > 0 else 0.0
                    row_cells[col_idx].text = f"{b_nota:.2f}"
                    col_idx += 1

            else: # view_mode == "disciplina"
                for d_name in mid_full:
                    if not is_subject_applicable_to_tipo(d_name, r.get("tipo_prova", "")):
                        row_cells[col_idx].text = "-"
                    else:
                        d_info = det.get(d_name)
                        if not d_info:
                            for k, v in det.items():
                                if k.strip().lower() == d_name.strip().lower():
                                    d_info = v
                                    break
                        if not d_info:
                            d_info = {}
                        tot = d_info.get("total", 0)
                        ac = d_info.get("acertos", 0)
                        n_disc = (ac / tot * 10.0) if tot > 0 else d_info.get("nota", 0.0)
                        row_cells[col_idx].text = f"{n_disc:.2f}"
                    col_idx += 1

            if possui_redacao:
                n_red = r.get("nota_redacao")
                n_red_str = f"{n_red:.2f}" if n_red is not None else "0.00"
                row_cells[col_idx].text = n_red_str
                col_idx += 1

            if possui_redacao or view_mode != "geral":
                n_red_val = r.get("nota_redacao") or 0.0 if possui_redacao else 0.0
                fn_val = float(r.get("total_acertos", 0)) + n_red_val
                row_cells[col_idx].text = f"{fn_val:.2f}"

        doc.save(output_path)
