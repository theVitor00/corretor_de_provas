import os
import pytest
import openpyxl
import docx
from PyQt6.QtWidgets import QApplication
from database import Database
from models.exam import ExamModel
from models.student import StudentModel
from models.processing import ProcessingModel
from models.subject import SubjectBlockModel
from services.grading_engine import GradingEngine
from services.exporter import ReportExporter
from ui.tabs.reports_tab import ReportsTab
from ui.tabs.exams_tab import ExamFormDialog

@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app

@pytest.fixture
def test_db(tmp_path):
    db_file = os.path.join(tmp_path, "test_enem.db")
    return Database(db_file)

def test_grading_engine_enem_calculations():
    # 1 Bloco de 45 questões
    gabarito_45 = "A" * 45
    exam_1 = {
        "valor_total": 1000.0,
        "tipo_prova": "Simulado Enem",
        "gabaritos": {"1": gabarito_45},
        "layout_config": {
            "subjects": [
                {"tipo": None, "nome": "Linguagens, Códigos e Suas Tecnologias", "start_q": 1, "end_q": 45, "peso": 1.0}
            ]
        }
    }

    engine_1 = GradingEngine(exam_1)

    # 45 acertos -> 1000 pontos
    res_100 = engine_1.grade_student("1", "A" * 45)
    assert res_100["total_acertos"] == 45
    det_lc = res_100["detalhes_disciplinas"]["Linguagens, Códigos e Suas Tecnologias"]
    assert det_lc["nota"] == 1000
    assert res_100["nota_final"] == 1000.0

    # 0 acertos -> 0 pontos
    res_0 = engine_1.grade_student("1", "B" * 45)
    assert res_0["total_acertos"] == 0
    assert res_0["detalhes_disciplinas"]["Linguagens, Códigos e Suas Tecnologias"]["nota"] == 0
    assert res_0["nota_final"] == 0.0

    # 30 acertos -> round((1000 / 45) * 30) = round(666.666...) = 667 pontos
    res_30 = engine_1.grade_student("1", ("A" * 30) + ("B" * 15))
    assert res_30["total_acertos"] == 30
    assert res_30["detalhes_disciplinas"]["Linguagens, Códigos e Suas Tecnologias"]["nota"] == 667
    assert res_30["nota_final"] == 667.0

    # 2 Blocos de 45 questões (90 questões total)
    exam_2 = {
        "valor_total": 1000.0,
        "tipo_prova": "Simulado Enem",
        "gabaritos": {"1": "A" * 90},
        "layout_config": {
            "subjects": [
                {"tipo": None, "nome": "Linguagens, Códigos e Suas Tecnologias", "start_q": 1, "end_q": 45, "peso": 1.0},
                {"tipo": None, "nome": "Ciências Humanas e Suas Tecnologias", "start_q": 46, "end_q": 90, "peso": 1.0}
            ]
        }
    }

    engine_2 = GradingEngine(exam_2)

    # Bloco 1: 30 acertos (667), Bloco 2: 45 acertos (1000)
    # Nota Final = média dos blocos = round((667 + 1000) / 2) = 834
    res_blocos = engine_2.grade_student("1", ("A" * 30) + ("B" * 15) + ("A" * 45))
    assert res_blocos["total_acertos"] == 75
    det_b1 = res_blocos["detalhes_disciplinas"]["Linguagens, Códigos e Suas Tecnologias"]
    det_b2 = res_blocos["detalhes_disciplinas"]["Ciências Humanas e Suas Tecnologias"]
    assert det_b1["nota"] == 667
    assert det_b2["nota"] == 1000
    assert res_blocos["nota_final"] == 834.0


def test_exam_form_enem_auto_align(qapp, test_db):
    exam_model = ExamModel(test_db)
    subject_model = SubjectBlockModel(test_db)
    dlg = ExamFormDialog(exam_model, subject_model)
    
    # Seleciona Simulado Enem
    idx = dlg.combo_tipo_prova.findText("Simulado Enem")
    assert idx >= 0
    dlg.combo_tipo_prova.setCurrentIndex(idx)

    # Inicia com 0 blocos
    assert len(dlg.selected_blocos) == 0
    assert dlg.spin_valor.value() == 1000.0

    # Adiciona Bloco 1: Linguagens, Códigos e Suas Tecnologias
    dlg.selected_blocos = [
        {"id": 1, "nome": "Linguagens, Códigos e Suas Tecnologias", "sigla": "LC"}
    ]
    dlg.auto_align_enem_blocks()
    
    assert len(dlg.mapped_subjects) == 1
    assert dlg.mapped_subjects[0]["nome"] == "Linguagens, Códigos e Suas Tecnologias"
    assert dlg.mapped_subjects[0]["start_q"] == 1
    assert dlg.mapped_subjects[0]["end_q"] == 45
    assert dlg.spin_valor.value() == 1000.0

    # Adiciona Bloco 2: Ciências Humanas e Suas Tecnologias
    dlg.selected_blocos.append(
        {"id": 2, "nome": "Ciências Humanas e Suas Tecnologias", "sigla": "CH"}
    )
    dlg.auto_align_enem_blocks()

    assert len(dlg.mapped_subjects) == 2
    assert dlg.mapped_subjects[0]["nome"] == "Linguagens, Códigos e Suas Tecnologias"
    assert dlg.mapped_subjects[0]["start_q"] == 1
    assert dlg.mapped_subjects[0]["end_q"] == 45
    assert dlg.mapped_subjects[1]["nome"] == "Ciências Humanas e Suas Tecnologias"
    assert dlg.mapped_subjects[1]["start_q"] == 46
    assert dlg.mapped_subjects[1]["end_q"] == 90
    assert dlg.spin_valor.value() == 1000.0

    # Testar reordenação (inverter ordem dos blocos)
    dlg.selected_blocos[0], dlg.selected_blocos[1] = dlg.selected_blocos[1], dlg.selected_blocos[0]
    dlg.auto_align_enem_blocks()

    assert dlg.mapped_subjects[0]["nome"] == "Ciências Humanas e Suas Tecnologias"
    assert dlg.mapped_subjects[0]["start_q"] == 1
    assert dlg.mapped_subjects[0]["end_q"] == 45
    assert dlg.mapped_subjects[1]["nome"] == "Linguagens, Códigos e Suas Tecnologias"
    assert dlg.mapped_subjects[1]["start_q"] == 46
    assert dlg.mapped_subjects[1]["end_q"] == 90


def test_reports_tab_enem_view_modes(qapp, test_db):
    exam_model = ExamModel(test_db)
    subject_model = SubjectBlockModel(test_db)
    proc_model = ProcessingModel(test_db)

    b1_id = subject_model.create_block("Linguagens, Códigos e Suas Tecnologias")
    b2_id = subject_model.create_block("Ciências Humanas e Suas Tecnologias")

    # Criar prova Simulado Enem com 2 blocos
    layout_config = {
        "partes": [{"parte_num": 1, "nome": "Parte Única", "num_questoes": 90, "q_start": 1, "q_end": 90}],
        "subjects": [
            {"tipo": None, "nome": "Linguagens, Códigos e Suas Tecnologias", "start_q": 1, "end_q": 45, "peso": 1.0},
            {"tipo": None, "nome": "Ciências Humanas e Suas Tecnologias", "start_q": 46, "end_q": 90, "peso": 1.0}
        ]
    }
    exam_id = exam_model.create_exam(
        nome="Simulado ENEM 2026",
        data="2026-10-20",
        tipo_prova="Simulado Enem",
        gabaritos={"1": "A" * 90},
        valor_total=1000.0,
        layout_config=layout_config,
        bloco_ids=[b1_id, b2_id],
        possui_redacao=True
    )
    exam = exam_model.get_exam_by_id(exam_id)

    # Inserir resultado de aluno com nota calculada de ENEM (média dos blocos = 834)
    detalhes = {
        "Linguagens, Códigos e Suas Tecnologias": {"acertos": 30, "total": 45, "percentual": 66.7, "nota": 667},
        "Ciências Humanas e Suas Tecnologias": {"acertos": 45, "total": 45, "percentual": 100.0, "nota": 1000}
    }
    proc_model.save_processing(
        prova_id=exam_id,
        aluno_matricula="1001",
        tipo_prova="1",
        respostas_aluno=("A" * 30) + ("B" * 15) + ("A" * 45),
        status_controle="OK",
        nota_final=834.0,
        nota_redacao=800.0,
        percentual_acertos=83.33,
        total_acertos=75,
        total_questoes=90,
        detalhes_disciplinas=detalhes
    )

    tab = ReportsTab()
    tab.exam_model = exam_model
    tab.subject_model = subject_model
    tab.processing_model = proc_model
    tab.current_exam = exam

    # 1. Verificar restrição do combo de modo de visualização
    tab._update_view_mode_combo_for_exam(exam)
    items = [tab.combo_view_mode.itemData(i) for i in range(tab.combo_view_mode.count())]
    assert items == ["bloco", "geral"]

    # 2. Testar carregamento da tabela no modo "bloco" (Notas)
    tab.select_exam(exam_id)
    tab.combo_view_mode.setCurrentIndex(tab.combo_view_mode.findData("bloco"))
    tab.load_results()

    # Verificar que as colunas preservam a ordem definida (LC depois CH, e NÃO ordem alfabética)
    # Índices: 0: Seq, 1: Matrícula, 2: Nome, 3: Turma, 4: Modelo, 5: LC, 6: CH, 7: Redação, 8: Nota Final, 9: Ação
    header_5 = tab.tbl_results.horizontalHeaderItem(5).text()
    header_6 = tab.tbl_results.horizontalHeaderItem(6).text()
    assert header_5 == "LC"
    assert header_6 == "CH"

    # Verificar valores inteiros
    val_lc = tab.tbl_results.item(0, 5).text()
    val_ch = tab.tbl_results.item(0, 6).text()
    val_red = tab.tbl_results.item(0, 7).text()
    val_final = tab.tbl_results.item(0, 8).text()

    assert val_lc == "667"
    assert val_ch == "1000"
    assert val_red == "800"
    # A Nota Final é a média dos blocos (834), desconsiderando a Redação (que fica como valor separado na coluna anterior)
    assert val_final == "834"


def test_exporter_enem_outputs(test_db, tmp_path):
    exam_model = ExamModel(test_db)
    subject_model = SubjectBlockModel(test_db)
    proc_model = ProcessingModel(test_db)
    exporter = ReportExporter(test_db)

    b1_id = subject_model.create_block("Linguagens, Códigos e Suas Tecnologias")
    b2_id = subject_model.create_block("Ciências Humanas e Suas Tecnologias")

    layout_config = {
        "partes": [{"parte_num": 1, "nome": "Parte Única", "num_questoes": 90, "q_start": 1, "q_end": 90}],
        "subjects": [
            {"tipo": None, "nome": "Linguagens, Códigos e Suas Tecnologias", "start_q": 1, "end_q": 45, "peso": 1.0},
            {"tipo": None, "nome": "Ciências Humanas e Suas Tecnologias", "start_q": 46, "end_q": 90, "peso": 1.0}
        ]
    }
    exam_id = exam_model.create_exam(
        nome="Simulado ENEM 2026",
        data="2026-10-20",
        tipo_prova="Simulado Enem",
        gabaritos={"1": "A" * 90},
        valor_total=1000.0,
        layout_config=layout_config,
        bloco_ids=[b1_id, b2_id],
        possui_redacao=True
    )
    exam = exam_model.get_exam_by_id(exam_id)

    detalhes = {
        "Linguagens, Códigos e Suas Tecnologias": {"acertos": 30, "total": 45, "percentual": 66.7, "nota": 667},
        "Ciências Humanas e Suas Tecnologias": {"acertos": 45, "total": 45, "percentual": 100.0, "nota": 1000}
    }
    proc_model.save_processing(
        prova_id=exam_id,
        aluno_matricula="1001",
        tipo_prova="1",
        respostas_aluno=("A" * 30) + ("B" * 15) + ("A" * 45),
        status_controle="OK",
        nota_final=834.0,
        nota_redacao=800.0,
        percentual_acertos=83.33,
        total_acertos=75,
        total_questoes=90,
        detalhes_disciplinas=detalhes
    )
    results = proc_model.list_results_for_exam(exam_id)

    # 1. Exportar Excel no modo "bloco"
    excel_path = str(tmp_path / "simulado_enem.xlsx")
    exporter.export_exam_excel(exam, results, excel_path, view_mode="bloco")
    assert os.path.exists(excel_path)

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    # Encontrar linha do aluno
    all_vals = []
    for r in ws.iter_rows(values_only=True):
        if 667 in r:
            all_vals = list(r)
            break
    assert 667 in all_vals
    assert 1000 in all_vals
    assert 800 in all_vals
    assert 834 in all_vals

    # 2. Exportar Word no modo "bloco"
    word_path = str(tmp_path / "simulado_enem.docx")
    exporter.export_exam_word(exam, results, word_path, view_mode="bloco")
    assert os.path.exists(word_path)

    doc = docx.Document(word_path)
    t = doc.tables[0]
    cell_texts = [cell.text for cell in t.rows[1].cells]
    assert "667" in cell_texts
    assert "1000" in cell_texts
    assert "800" in cell_texts
    assert "834" in cell_texts

    # 3. Exportar PDF no modo "bloco"
    pdf_path = str(tmp_path / "simulado_enem.pdf")
    exporter.export_exam_pdf(exam, results, pdf_path, view_mode="bloco")
    assert os.path.exists(pdf_path)
    assert os.path.getsize(pdf_path) > 1000
