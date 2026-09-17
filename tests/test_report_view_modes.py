import os
import pytest
import openpyxl
from PyQt6.QtWidgets import QApplication, QPushButton, QTableWidget
from database import Database
from models.exam import ExamModel
from models.student import StudentModel
from models.processing import ProcessingModel
from models.subject import SubjectBlockModel
from services.exporter import ReportExporter
from ui.tabs.reports_tab import ReportsTab

@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app

@pytest.fixture
def test_db(tmp_path):
    db_file = os.path.join(tmp_path, "test_corretor.db")
    return Database(db_file)

def test_export_and_ui_data_integrity(qapp, test_db, tmp_path):
    exam_model = ExamModel(test_db)
    student_model = StudentModel(test_db)
    proc_model = ProcessingModel(test_db)
    subject_model = SubjectBlockModel(test_db)
    exporter = ReportExporter(test_db)

    # 1. Criar Prova com múltiplas disciplinas
    gabaritos = {"1": "A" * 45}
    layout_config = {
        "partes": [{"parte_num": 1, "nome": "Parte Única", "num_questoes": 45, "q_start": 1, "q_end": 45}],
        "subjects": [
            {"tipo": "1", "nome": "Português", "start_q": 1, "end_q": 10},
            {"tipo": "1", "nome": "Matemática", "start_q": 11, "end_q": 20},
            {"tipo": "1", "nome": "Biologia", "start_q": 21, "end_q": 30},
            {"tipo": "1", "nome": "História", "start_q": 31, "end_q": 40},
            {"tipo": "1", "nome": "Física", "start_q": 41, "end_q": 45},
        ]
    }
    exam_id = exam_model.create_exam(
        nome="Prova Completa Integridade",
        data="2026-09-15",
        gabaritos=gabaritos,
        valor_total=10.0,
        layout_config=layout_config
    )
    exam = exam_model.get_exam_by_id(exam_id)

    # 2. Inserir Alunos Processados
    detalhes_a1 = {
        "Português": {"acertos": 8, "total": 10, "percentual": 80.0, "nota": 8.0},
        "Matemática": {"acertos": 7, "total": 10, "percentual": 70.0, "nota": 7.0},
        "Biologia": {"acertos": 9, "total": 10, "percentual": 90.0, "nota": 9.0},
        "História": {"acertos": 10, "total": 10, "percentual": 100.0, "nota": 10.0},
        "Física": {"acertos": 4, "total": 5, "percentual": 80.0, "nota": 8.0},
    }
    proc_model.save_processing(
        prova_id=exam_id,
        aluno_matricula="201",
        tipo_prova="1",
        respostas_aluno="A" * 38 + "B" * 7,
        status_controle="OK",
        nota_final=8.44,
        percentual_acertos=84.44,
        total_acertos=38,
        total_questoes=45,
        detalhes_disciplinas=detalhes_a1
    )

    results = proc_model.list_results_for_exam(exam_id)
    assert len(results) == 1

    # 3. Testar Integridade na Interface (ReportsTab) - DEVE conter a coluna Ação na GUI
    tab = ReportsTab()
    tab.exam_model = exam_model
    tab.student_model = student_model
    tab.processing_model = proc_model
    tab.subject_model = subject_model

    tab.select_exam(exam_id)

    for mode in ["geral", "bloco", "disciplina"]:
        idx = tab.combo_view_mode.findData(mode)
        tab.combo_view_mode.setCurrentIndex(idx)
        tab.load_results()

        tbl = tab.tbl_results
        total_cols = tbl.columnCount()
        action_col = total_cols - 1

        # A última coluna da interface GUI DEVE ser a coluna 'Ação'
        header_action_item = tbl.horizontalHeaderItem(action_col)
        assert header_action_item is not None
        assert header_action_item.text() == "Ação"

        # Apenas a coluna Ação na GUI deve conter o botão
        for row in range(tbl.rowCount()):
            for col in range(total_cols):
                widget = tbl.cellWidget(row, col)
                if col == action_col:
                    assert isinstance(widget, QPushButton)
                    assert widget.text() == "Boletim Individual"
                else:
                    assert widget is None

            if mode == "disciplina":
                bio_col_idx = 7
                bio_item = tbl.item(row, bio_col_idx)
                assert bio_item is not None and bio_item.text() == "9.00"

                his_col_idx = 8
                his_item = tbl.item(row, his_col_idx)
                assert his_item is not None and his_item.text() == "10.00"

    # 4. Testar Exportação para Arquivos (PDF, Excel, Word) - NÃO DEVE conter a coluna Ação!
    for mode in ["geral", "bloco", "disciplina"]:
        pdf_path = str(tmp_path / f"test_integ_{mode}.pdf")
        exporter.export_exam_pdf(exam, results, pdf_path, view_mode=mode)
        assert os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0

        xlsx_path = str(tmp_path / f"test_integ_{mode}.xlsx")
        exporter.export_exam_excel(exam, results, xlsx_path, view_mode=mode)
        assert os.path.exists(xlsx_path) and os.path.getsize(xlsx_path) > 0

        # Verificar se a planilha Excel exportada NÃO possui a coluna 'Ação'
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        headers_row5 = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Ação" not in headers_row5, f"A coluna 'Ação' não deve estar presente no Excel exportado ({mode})"

        docx_path = str(tmp_path / f"test_integ_{mode}.docx")
        exporter.export_exam_word(exam, results, docx_path, view_mode=mode)
        assert os.path.exists(docx_path) and os.path.getsize(docx_path) > 0
